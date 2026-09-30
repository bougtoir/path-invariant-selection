"""Build SHE submission outputs: docx (journal layout + blinded + inline-figure reading copy),
figures.pptx, cover letter, references.bib, and the submission zip."""
import os, re, zipfile, subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PIn, Pt as PPt

ROOT = os.path.join(os.path.dirname(__file__), "..")
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "output")
os.makedirs(OUT, exist_ok=True)

SRC = open(os.path.join(ROOT, "manuscript_source.md")).read()

TITLE = "When Should Experience Matter? Path-Invariant Selection and the Credentialisation of Experience in Higher Education"

# ---------- parse manuscript_source.md ----------
def parse(src):
    sections = []  # (level, heading, [paragraphs])
    cur = None
    for block in src.split("\n\n"):
        b = block.strip()
        if not b or b.startswith("---"):
            continue
        if b.startswith("## "):
            cur = (b[3:].strip(), [])
            sections.append(cur)
        elif b.startswith("# "):
            continue
        else:
            if cur is None:
                cur = ("__front__", [])
                sections.append(cur)
            for line in b.split("\n"):
                line = line.strip()
                if line:
                    cur[1].append(line)
    return sections

sections = parse(SRC)

def wc(text):
    return len(re.findall(r"[A-Za-z0-9'-]+", text))

main_words = 0
for h, paras in sections:
    if h in ("__front__",):
        continue
    if h in ("References", "Disclosure statement", "Funding", "Acknowledgements"):
        continue
    for p in paras:
        main_words += wc(p)
    main_words += wc(h)

ref_words = 0
for h, paras in sections:
    if h == "References":
        ref_words = sum(wc(p) for p in paras)
print("main-text words:", main_words, "| reference words:", ref_words, "| est. total:", main_words + ref_words + 500)

MD_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")

def add_para(doc, text, style=None, align=None):
    p = doc.add_paragraph(style=style)
    for chunk in MD_RE.split(text):
        if chunk.startswith("**") and chunk.endswith("**"):
            r = p.add_run(chunk[2:-2]); r.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            r = p.add_run(chunk[1:-1]); r.italic = True
        else:
            p.add_run(chunk)
    if align:
        p.alignment = align
    return p

def set_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)

def build_docx(blinded=False, inline_figs=False, path=""):
    doc = Document()
    set_style(doc)
    fig_cites = {"Figure 1": "Figure1.png", "Figure 2": "Figure2.png", "Figure 3": "Figure3.png"}
    fig_done = set()

    if not blinded:
        tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(TITLE); r.bold = True; r.font.size = Pt(14)
        doc.add_paragraph()
        for line in [
            "Title page",
            "",
            "Author: [AUTHOR NAME — to be completed]",
            "Affiliation: [AFFILIATION — to be completed]",
            "Postal address: [ADDRESS — to be completed]",
            "Email / telephone: [CONTACT — to be completed]",
            "Corresponding author: [YES/NO]",
            "",
            f"Word count (main text, excl. references, tables, captions): approx. {main_words}",
            "Acknowledgements: [to be completed by the author]",
            "Funding: [to be completed by the author]",
            "Disclosure statement: No potential conflict of interest was reported by the author(s).",
        ]:
            add_para(doc, line)
        doc.add_page_break()

    if blinded:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(14)
        cap = doc.add_paragraph("Manuscript prepared for double-anonymised review. No author-identifying information is included.")

    skip_heads = {"__front__"}
    for h, paras in sections:
        if h in skip_heads:
            continue
        if blinded and h in ("Acknowledgements", "Funding"):
            continue
        if h in ("Abstract", "Keywords"):
            doc.add_heading(h, level=1)
            for p in paras:
                add_para(doc, p)
            continue
        doc.add_heading(h, level=1)
        for p in paras:
            if p.startswith("- "):
                for item in p.split("\n"):
                    pass
                add_para(doc, p[2:] if p.startswith("- ") else p, style="List Bullet")
            else:
                add_para(doc, p)
            if inline_figs:
                for label, fn in fig_cites.items():
                    if label in p and label not in fig_done:
                        fig_done.add(label)
                        doc.add_picture(os.path.join(FIG, fn), width=Inches(6.0))
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        rc = cp.add_run(f"{label}. See caption list.")
                        rc.font.size = Pt(10); rc.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # tables, one per page
    for tname, tfile, ttitle in [
        ("Table 1", "tables/Table1_three_selection_environments.md",
         "Table 1. Three selection environments."),
        ("Table 2", "tables/Table2_relation_to_theory.md",
         "Table 2. Relation of the path-invariance framework to existing theory."),
    ]:
        doc.add_page_break()
        add_para(doc, ttitle).runs[0].bold = True
        rows = [l.strip() for l in open(os.path.join(ROOT, tfile)) if l.strip().startswith("|")]
        header = [c.strip() for c in rows[0].strip("|").split("|")]
        body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
        tbl = doc.add_table(rows=1, cols=len(header))
        tbl.style = "Light Grid Accent 1"
        for i, c in enumerate(header):
            tbl.rows[0].cells[i].text = c
            for rr in tbl.rows[0].cells[i].paragraphs[0].runs:
                rr.bold = True; rr.font.size = Pt(9)
        for row in body:
            cells = tbl.add_row().cells
            for i, c in enumerate(row[:len(header)]):
                cells[i].text = c
                for para in cells[i].paragraphs:
                    for rr in para.runs:
                        rr.font.size = Pt(9)
        if tname == "Table 1":
            add_para(doc, "Note. Positions on the direct-measurement continuum are conceptual and illustrative, not empirically estimated.")

    # figure captions
    doc.add_page_break()
    doc.add_heading("Figure captions", level=1)
    caps = [
        "Figure 1. The two pathways from experience to selection. Upper path: resources enable experience, which builds competence, which produces later outcomes (developmental). Lower path: experience generates a credential signal that enters evaluation directly, independently of competence (credential).",
        "Figure 2. Path-invariant versus path-dependent selection. Panel A: differing acquisition pathways that produce equivalent relevant capability receive equivalent evaluations. Panel B: the pathway itself shifts evaluation despite equivalent capability — a violation requiring justification.",
        "Figure 3. The direct-measurement continuum. Conceptual, illustrative positioning of three selection environments: where valid direct measurement is less feasible, residual reliance on biography is expected to be greater. Positions are not empirically estimated.",
    ]
    for c in caps:
        add_para(doc, c)

    doc.save(path)
    return path

final = build_docx(path=os.path.join(OUT, "SHE_manuscript_final.docx"))
blind = build_docx(blinded=True, path=os.path.join(OUT, "SHE_manuscript_blinded.docx"))
inlin = build_docx(inline_figs=True, path=os.path.join(OUT, "SHE_manuscript_inline_figures.docx"))

# ---------- figures.pptx ----------
prs = Presentation()
prs.slide_width = PIn(10); prs.slide_height = PIn(7.5)
titles = {
    "Figure1.png": "Figure 1. The two pathways from experience to selection",
    "Figure2.png": "Figure 2. Path-invariant versus path-dependent selection",
    "Figure3.png": "Figure 3. The direct-measurement continuum",
}
for fn, t in titles.items():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(PIn(0.3), PIn(0.2), PIn(9.4), PIn(0.7))
    tb.text_frame.text = t
    tb.text_frame.paragraphs[0].runs[0].font.size = PPt(16)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    s.shapes.add_picture(os.path.join(FIG, fn), PIn(0.8), PIn(1.0), width=PIn(8.4))
prs.save(os.path.join(OUT, "figures.pptx"))

# ---------- cover letter ----------
cl = Document()
set_style(cl)
add_para(cl, "Cover letter — Studies in Higher Education").runs[0].bold = True
for para in [
    "",
    "Dear Editors,",
    "",
    "We submit the enclosed manuscript, \u2018When Should Experience Matter? Path-Invariant Selection and the Credentialisation of Experience in Higher Education\u2019, for consideration in Studies in Higher Education.",
    "",
    "The manuscript addresses a problem at the centre of current higher-education admissions practice: universities increasingly evaluate applicants through experiences, activities and biographies, yet the literature lacks a criterion that distinguishes experience as a means of becoming capable from experience as a fact independently rewarded. The article develops the path-invariance principle — where acquisition pathways are substitutable with respect to the attributes an institution legitimately values, equivalent demonstrated attributes should receive equivalent evaluation — and derives a framework (the development/credential distinction, substitutability, and the direct-measurement continuum) together with six conditional propositions and a compact, falsifiable empirical agenda.",
    "",
    "University admissions is the paper's central case and carries the argument. Employment selection and political candidacy are used only to establish the generality of the selection architecture and to mark its boundary where direct measurement is least feasible. The manuscript is not advocacy for any particular admissions technology: it explicitly protects contextualised admissions and does not claim that standardised tests eliminate inequality.",
    "",
    "We believe the article will interest the journal's readership across admissions research, higher-education sociology, and selection policy. The manuscript is original, is not under consideration elsewhere, and complies with the journal's formatting, anonymisation and length guidance.",
    "",
    "Yours sincerely,",
    "",
    "[Author name, affiliation, contact — to be completed]",
]:
    add_para(cl, para)
cl.save(os.path.join(OUT, "cover_letter_SHE.docx"))

# ---------- references.bib ----------
bib = open(os.path.join(ROOT, "refs", "references.bib"), "w")
bib.write(re.sub(r"^.*$", "", "") or "")
bib.close()
# generate a simple bib from the reference list
refs = [p for h, paras in sections if h == "References" for p in paras]
with open(os.path.join(ROOT, "refs", "references.bib"), "w") as f:
    for i, p in enumerate(refs):
        key = re.sub(r"[^A-Za-z]", "", p.split(",")[0]) + re.search(r"\b(19|20)\d{2}\b", p).group(0) if re.search(r"\b(19|20)\d{2}\b", p) else f"ref{i}"
        doi = re.search(r"doi:(\S+)", p)
        f.write(f"@misc{{{key},\n  note = {{{p}}}" + (f",\n  doi = {{{doi.group(1).rstrip('.')}}}" if doi else "") + "\n}\n\n")

# ---------- PDF ----------
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT,
                os.path.join(OUT, "SHE_manuscript_final.docx")], check=True, capture_output=True)

# ---------- copy figures ----------
import shutil
for fn in os.listdir(FIG):
    if fn.endswith(".png"):
        shutil.copy(os.path.join(FIG, fn), os.path.join(OUT, fn))

# ---------- zip ----------
zpath = os.path.join(OUT, "SHE_submission_package_FINAL.zip")
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for fn in sorted(os.listdir(OUT)):
        if fn.endswith(".zip"):
            continue
        z.write(os.path.join(OUT, fn), fn)
    for extra in ["manuscript_source.md", "NOVELTY_ASSESSMENT.md", "SHE_FIT_AUDIT.md",
                  "HOSTILE_REVIEW.md", "DESK_REJECTION_AUDIT.md", "REFERENCE_AUDIT.md",
                  "LITERATURE_MAP.md", "SHE_AUTHOR_GUIDE.md", "THEORY_SPECIFICATION.md",
                  "PROJECT_SCOPE.md", "EXECUTION_LOG.md"]:
        z.write(os.path.join(ROOT, extra), extra)
    z.write(os.path.join(ROOT, "refs", "references.bib"), "references.bib")
print("built:", sorted(os.listdir(OUT)))
